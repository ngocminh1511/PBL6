r"""Build the v2 retrieval test set for the whole Lao dong topic (chu de 20).

The v1 benchmark has 30 questions, all on de muc 20.5 and almost all on the
Luat.  v2 covers 20.2 / 20.3 / 20.4 / 20.5, several question styles, questions
whose answer sits in a late Khoan of a long Dieu, and questions that need two
Dieu.  Labels are LLM-assisted but human-verified:

  sample    stratified seed Dieu per de muc / document type / style  -> seeds.jsonl
  draft     LLM writes one question per seed (resumable)             -> drafts.jsonl
  pool      candidate Dieu per question from BM25 + 3 embedding
            models (IR "pooling"), so labels are not biased toward
            one retriever                                            -> pools.jsonl
  judge     LLM suggests a 0-3 grade per candidate (resumable)       -> judgments.jsonl
  export    Excel workbook for human review; re-running it keeps
            every value a reviewer has already typed                 -> review.xlsx
  finalize  reviewed workbook -> benchmark_questions_v2.jsonl
  stats     coverage report                                          -> testset_stats.json

The 30 v1 questions are carried over (ids q01..q30) and re-pooled against the
full topic, because Dieu outside 20.5 (e.g. the Bo luat Lao dong chapter on
an toan lao dong) may now also answer them.

Run from the repository root:

    python -m labor_rag.evaluation.build_testset sample
    python -m labor_rag.evaluation.build_testset draft
    python -m labor_rag.evaluation.build_testset pool
    python -m labor_rag.evaluation.build_testset judge
    python -m labor_rag.evaluation.build_testset export
    # ... review labor_rag/evaluation/annotation/review.xlsx (see GUIDELINES.md) ...
    python -m labor_rag.evaluation.build_testset finalize
    python -m labor_rag.evaluation.build_testset stats
"""
from __future__ import annotations

import argparse
import hashlib
import json
import math
import random
import re
import sys
import time
import unicodedata
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any, Iterable

from labor_rag.config import ANNOTATION_DIR, EVALUATION_DIR, LAO_DONG_V4_DIR, TESTSET_CACHE_DIR

DEFAULT_CHUNKS = LAO_DONG_V4_DIR / "chunks.jsonl"
DEFAULT_EDGES = LAO_DONG_V4_DIR / "kg_edges.jsonl"
V1_QUESTIONS = EVALUATION_DIR / "benchmark_questions.jsonl"
OUTPUT_QUESTIONS = EVALUATION_DIR / "benchmark_questions_v2.jsonl"

SEEDS_PATH = ANNOTATION_DIR / "seeds.jsonl"
DRAFTS_PATH = ANNOTATION_DIR / "drafts.jsonl"
POOLS_PATH = ANNOTATION_DIR / "pools.jsonl"
JUDGMENTS_PATH = ANNOTATION_DIR / "judgments.jsonl"
REVIEW_PATH = ANNOTATION_DIR / "review.xlsx"
STATS_PATH = ANNOTATION_DIR / "testset_stats.json"

# New questions per de muc (the 30 v1 questions add to 20.5 on top of this).
DEFAULT_QUOTAS = {"20.2": 50, "20.3": 20, "20.4": 30, "20.5": 20}
STYLE_RATIOS = [
    ("tinh_huong", 0.35),
    ("doi_thuong", 0.25),
    ("truc_tiep", 0.20),
    ("tu_khoa", 0.10),
    ("nhieu_dieu", 0.10),
]
DEEP_RATIO = 0.25  # share of a de muc quota asked about a late Khoan of a long Dieu
EMBEDDING_TOKEN_LIMIT = 256  # bkai max_seq_length: text past this is never embedded

CATEGORIES = [
    "hop_dong_lao_dong",
    "tien_luong",
    "thoi_gio_lam_viec_nghi_ngoi",
    "ky_luat_trach_nhiem_vat_chat",
    "cham_dut_hop_dong",
    "lao_dong_dac_thu",
    "doi_thoai_thuong_luong_cong_doan",
    "tranh_chap_lao_dong",
    "viec_lam_ho_tro_viec_lam",
    "bao_hiem_that_nghiep",
    "di_lam_viec_o_nuoc_ngoai",
    "an_toan_ve_sinh_lao_dong",
    "tai_nan_benh_nghe_nghiep",
    "quan_ly_nha_nuoc_thanh_tra_xu_ly",
    "khac",
]

STYLE_INSTRUCTIONS = {
    "truc_tiep": "Câu hỏi trực tiếp, rõ ràng về nội dung quy định (nhưng KHÔNG lặp lại nguyên văn tiêu đề Điều).",
    "tinh_huong": (
        "Câu hỏi tình huống thực tế ở ngôi thứ nhất (\"Tôi...\", \"Công ty tôi...\"), 1-3 câu ngắn, "
        "mô tả hoàn cảnh cụ thể rồi hỏi điều luật quy định thế nào."
    ),
    "doi_thuong": (
        "Câu hỏi bằng ngôn ngữ đời thường của người lao động bình thường, KHÔNG dùng thuật ngữ pháp lý "
        "có trong đoạn luật (ví dụ nói \"bị đuổi việc\" thay vì \"đơn phương chấm dứt hợp đồng\")."
    ),
    "tu_khoa": "Một cụm từ khóa tìm kiếm ngắn (3-8 từ), như người dùng gõ vào ô tìm kiếm, không phải câu hoàn chỉnh.",
    "nhieu_dieu": (
        "Một câu hỏi mà muốn trả lời đầy đủ thì PHẢI dùng thông tin từ CẢ HAI đoạn luật bên dưới "
        "(ví dụ điều kiện ở đoạn 1 và mức hưởng/thủ tục ở đoạn 2). Có thể là tình huống."
    ),
}

MAPC_RE = re.compile(r"^20\.(\d+)\.([A-ZĐ]+)\.[0-9][0-9.]*[a-zđ]*$")
BOILERPLATE_RE = re.compile(
    r"^(phạm vi điều chỉnh|đối tượng áp dụng|hiệu lực( thi hành)?|trách nhiệm thi hành|"
    r"tổ chức thực hiện|điều khoản (thi hành|chuyển tiếp)|quy định chuyển tiếp|bãi bỏ|"
    r"sửa đổi, bổ sung)",
    re.IGNORECASE,
)
MIN_BODY_CHARS = 120
# Closing clauses whose title is only the Dieu code ("Điều 20.5.TT.58.3."), so
# the title check above misses them; recognise them by their body instead.
BOILERPLATE_BODY_RE = re.compile(
    r"chịu trách nhiệm thi hành|có hiệu lực (thi hành )?kể từ ngày|"
    r"văn bản được dẫn chiếu .{0,60}(thay thế|sửa đổi)",
    re.IGNORECASE,
)
MAX_BOILERPLATE_BODY_CHARS = 400
# "Chinh phu quy dinh chi tiet ..." Khoan only delegate; nothing to ask about.
DELEGATION_RE = re.compile(
    r"^\d+\.\s*(Chính phủ|Thủ tướng|Bộ trưởng|Bộ Lao động)[^.:]{0,120}?quy định (chi tiết|cụ thể|mẫu|việc)",
    re.IGNORECASE,
)
TOKEN_RE = re.compile(r"\w+", re.UNICODE)
ARTICLE_MENTION_RE = re.compile(r"Điều\s+\d|20\.\d+\.[A-ZĐ]+\.", re.IGNORECASE)

DENSE_MODELS = {
    "bkai": "bkai-foundation-models/vietnamese-bi-encoder",
    "bge": "BAAI/bge-m3",
    "mpnet": "sentence-transformers/paraphrase-multilingual-mpnet-base-v2",
}
TOKENIZER_MODEL = DENSE_MODELS["bkai"]


# --------------------------------------------------------------------------- io


def read_jsonl(path: Path) -> list[dict[str, Any]]:
    if not path.exists():
        return []
    return [json.loads(line) for line in path.read_text(encoding="utf-8").splitlines() if line.strip()]


def write_jsonl(path: Path, rows: Iterable[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="\n") as handle:
        for row in rows:
            handle.write(json.dumps(row, ensure_ascii=False) + "\n")


def append_jsonl(path: Path, row: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("a", encoding="utf-8", newline="\n") as handle:
        handle.write(json.dumps(row, ensure_ascii=False) + "\n")


def latest_by_qid(rows: list[dict[str, Any]]) -> dict[str, dict[str, Any]]:
    """Resumable files are append-only; the last record of a qid wins."""

    return {row["qid"]: row for row in rows}


# ------------------------------------------------------------------- articles


def de_muc_of(mapc: str) -> str:
    return ".".join(mapc.split(".")[:2])


def doc_type_of(mapc: str) -> str:
    parts = mapc.split(".")
    return parts[2] if len(parts) > 2 else "?"


def article_body(chunk: dict[str, Any]) -> str:
    """Chunk text is 'breadcrumb / title / citation / body'; return the body."""

    lines = chunk.get("text", "").split("\n")
    citation = (chunk.get("citation") or "").strip()
    for index, line in enumerate(lines):
        if citation and line.strip() == citation:
            return "\n".join(lines[index + 1 :]).strip()
    return chunk.get("text", "").strip()


def split_khoan(body: str) -> list[dict[str, Any]]:
    """Split a Dieu body into Khoan by looking for "1. ", "2. ", ... in order.

    Searching for the numbers sequentially avoids treating "khoản 2 Điều 35"
    or years as Khoan boundaries.  Returns [] when fewer than 2 Khoan exist.
    """

    starts: list[tuple[int, int]] = []
    position, number = 0, 1
    while True:
        match = re.compile(rf"(?:^|(?<=\s)){number}\.\s").search(body, position)
        if not match:
            break
        starts.append((number, match.start()))
        position, number = match.end(), number + 1
    if len(starts) < 2:
        return []
    khoan = []
    for index, (number, start) in enumerate(starts):
        end = starts[index + 1][1] if index + 1 < len(starts) else len(body)
        khoan.append({"khoan": number, "start": start, "text": body[start:end].strip()})
    return khoan


def is_eligible(chunk: dict[str, Any]) -> bool:
    mapc = str(chunk.get("mapc", ""))
    if not MAPC_RE.match(mapc):
        return False  # duplicate-code Dieu carry a URI as mapc; skip them as seeds
    if BOILERPLATE_RE.match((chunk.get("ten") or "").strip()):
        return False
    body = article_body(chunk)
    if len(body) <= MAX_BOILERPLATE_BODY_CHARS and BOILERPLATE_BODY_RE.search(body):
        return False
    return len(body) >= MIN_BODY_CHARS


def normalize(text: str) -> str:
    text = unicodedata.normalize("NFC", text or "").lower()
    return " ".join(TOKEN_RE.findall(text))


def tokens(text: str) -> list[str]:
    return TOKEN_RE.findall(unicodedata.normalize("NFC", text or "").lower())


class TokenCounter:
    """Counts tokens with the bkai tokenizer (the 256-token limit that matters)."""

    def __init__(self) -> None:
        from transformers import AutoTokenizer

        self.tokenizer = AutoTokenizer.from_pretrained(TOKENIZER_MODEL)

    def __call__(self, text: str) -> int:
        return len(self.tokenizer(text, add_special_tokens=True, verbose=False)["input_ids"])


def load_chunks(path: Path) -> dict[str, dict[str, Any]]:
    chunks = read_jsonl(path)
    if not chunks:
        raise SystemExit(f"No chunks at {path}. Build them first (see docs in this module).")
    return {str(chunk["mapc"]): chunk for chunk in chunks}


# ------------------------------------------------------------------------ llm


class QuotaExhausted(RuntimeError):
    """Daily (not per-minute) quota is used up; waiting/retrying is pointless."""


DAILY_QUOTA_RE = re.compile(r"PerDay|retry in \d+h", re.IGNORECASE)


def llm_call(system: str, user: str, max_tokens: int, retries: int = 5) -> str:
    """generation.llm_client.complete with backoff for free-tier rate limits."""

    from labor_rag.generation.llm_client import complete

    delay = 20.0
    for attempt in range(retries):
        try:
            return complete(system, user, max_tokens=max_tokens) or ""
        except Exception as exc:  # provider SDKs raise different types
            message = str(exc)
            if DAILY_QUOTA_RE.search(message):
                raise QuotaExhausted(message[:300]) from exc
            transient = any(code in message for code in ("429", "RESOURCE_EXHAUSTED", "quota", "503", "UNAVAILABLE", "overloaded"))
            if not transient or attempt == retries - 1:
                raise
            print(f"    rate limited/unavailable, retry in {delay:.0f}s ({message[:80]})")
            time.sleep(delay)
            delay *= 1.6
    return ""


GRADE_ITEM_RE = re.compile(r'\{\s*"i"\s*:\s*(\d+)\s*,\s*"grade"\s*:\s*([0-3])(?:\s*,\s*"why"\s*:\s*"([^"]*)")?')


def salvage_grades(raw: str) -> list[dict[str, Any]]:
    """Recover the complete items of a grades array cut off by max_tokens."""

    return [{"i": int(i), "grade": int(g), "why": why or ""} for i, g, why in GRADE_ITEM_RE.findall(raw)]


def stop_on_quota(exc: QuotaExhausted, path: Path) -> int:
    print(f"\nDaily LLM quota exhausted -- stopping. Progress so far is saved in {path.name};\n"
          "re-run the same command after the quota resets (or switch GEMINI_MODEL / LLM_PROVIDER in .env).\n"
          f"  {str(exc)[:160]}")
    return 2


def parse_json(raw: str) -> Any:
    """Take the first JSON object/array out of an LLM reply (tolerates ```json fences)."""

    raw = raw.strip()
    for opener, closer in (("{", "}"), ("[", "]")):
        start, end = raw.find(opener), raw.rfind(closer)
        if start != -1 and end > start:
            try:
                return json.loads(raw[start : end + 1])
            except json.JSONDecodeError:
                continue
    raise ValueError(f"No JSON in LLM reply: {raw[:200]!r}")


# --------------------------------------------------------------------- sample


def largest_remainder(weights: dict[str, float], total: int) -> dict[str, int]:
    weight_sum = sum(weights.values()) or 1.0
    exact = {key: total * value / weight_sum for key, value in weights.items()}
    result = {key: math.floor(value) for key, value in exact.items()}
    leftover = total - sum(result.values())
    for key in sorted(exact, key=lambda k: exact[k] - result[k], reverse=True)[:leftover]:
        result[key] += 1
    return result


def deep_target(chunk: dict[str, Any], count_tokens: TokenCounter) -> dict[str, Any] | None:
    """Pick a Khoan that starts past the embedding token limit (else in the
    second half of the Dieu) -- the content the current index cannot see."""

    text = chunk["text"]
    body = article_body(chunk)
    khoan = split_khoan(body)
    if len(khoan) < 3 or count_tokens(text) <= EMBEDDING_TOKEN_LIMIT:
        return None
    body_offset = text.find(body)
    beyond = [k for k in khoan if count_tokens(text[: body_offset + k["start"]]) > EMBEDDING_TOKEN_LIMIT]
    candidates = beyond or khoan[len(khoan) // 2 :]
    candidates = [k for k in candidates if len(k["text"]) >= 80 and not DELEGATION_RE.match(k["text"])]
    return random.choice(candidates) if candidates else None


def reference_pairs(edges_path: Path, eligible: set[str]) -> list[tuple[str, str]]:
    """Dieu -> Dieu citations inside the topic (e.g. a Nghi dinh citing the Luat)."""

    pairs = set()
    for edge in read_jsonl(edges_path):
        if edge.get("type") != "REFERS_TO":
            continue
        source, target = str(edge["source"]), str(edge["target"])
        if source in eligible and target in eligible and source != target:
            pairs.add((source, target))
    return sorted(pairs)


def cmd_sample(args: argparse.Namespace) -> int:
    if SEEDS_PATH.exists() and not args.force:
        raise SystemExit(f"{SEEDS_PATH} exists (drafts depend on it). Use --force to resample.")
    random.seed(args.seed)
    chunks = load_chunks(args.chunks)
    v1_gold = {mapc for question in read_jsonl(V1_QUESTIONS) for mapc in question["relevance"]}
    eligible = {mapc for mapc, chunk in chunks.items() if is_eligible(chunk)}
    count_tokens = TokenCounter()
    pairs = reference_pairs(args.edges, eligible)
    quotas = parse_quotas(args.quota) if args.quota else DEFAULT_QUOTAS

    used: set[str] = set(v1_gold)  # prefer Dieu the v1 set does not cover yet
    seeds: list[dict[str, Any]] = []
    for de_muc, quota in quotas.items():
        pool = sorted(m for m in eligible if de_muc_of(m) == de_muc and m not in used)
        if not pool:
            print(f"{de_muc}: no eligible Dieu, skipped")
            continue
        styles = [s for s, n in largest_remainder(dict(STYLE_RATIOS), quota).items() for _ in range(n)]
        random.shuffle(styles)
        n_multi = styles.count("nhieu_dieu")
        singles = [s for s in styles if s != "nhieu_dieu"]

        # 1) questions needing two Dieu, from real cross references; prefer
        #    pairs across document types (Nghi dinh/Thong tu -> Luat).
        local_pairs = [p for p in pairs if p[0] in pool and p[1] not in used]
        random.shuffle(local_pairs)
        local_pairs.sort(key=lambda p: doc_type_of(p[0]) == doc_type_of(p[1]))
        multi = []
        for source, target in local_pairs:
            if len(multi) == n_multi:
                break
            if source in used or target in used:
                continue
            multi.append((source, target))
            used.update((source, target))
        singles += ["tinh_huong"] * (n_multi - len(multi))  # not enough pairs: fall back

        # 2) deep-Khoan questions on long Dieu (never keyword style)
        n_deep = min(round(quota * DEEP_RATIO), sum(s != "tu_khoa" for s in singles))
        deep: list[tuple[str, dict[str, Any]]] = []
        long_pool = [m for m in pool if m not in used]
        random.shuffle(long_pool)
        for mapc in long_pool:
            if len(deep) == n_deep:
                break
            target = deep_target(chunks[mapc], count_tokens)
            if target:
                deep.append((mapc, target))
                used.add(mapc)

        # 3) remaining single-Dieu seeds, stratified by document type
        remaining = len(singles) - len(deep)
        rest_pool = [m for m in pool if m not in used]
        by_type: dict[str, list[str]] = defaultdict(list)
        for mapc in rest_pool:
            by_type[doc_type_of(mapc)].append(mapc)
        allocation = largest_remainder({t: len(v) for t, v in by_type.items()}, remaining)
        for doc_type, members in by_type.items():  # at least one per sizeable type
            if allocation[doc_type] == 0 and len(members) >= 5:
                donor = max(allocation, key=allocation.get)
                if allocation[donor] > 1:
                    allocation[donor] -= 1
                    allocation[doc_type] = 1
        plain: list[str] = []
        for doc_type, members in sorted(by_type.items()):
            plain += random.sample(members, min(allocation[doc_type], len(members)))
        used.update(plain)

        deep_styles = [s for s in singles if s != "tu_khoa"][: len(deep)]
        plain_styles = list(singles)
        for style in deep_styles:
            plain_styles.remove(style)
        for mapc, target in deep:
            seeds.append(make_seed(de_muc, deep_styles.pop(), [mapc], target))
        for mapc, style in zip(plain, plain_styles):
            seeds.append(make_seed(de_muc, style, [mapc], None))
        for source, target in multi:
            seeds.append(make_seed(de_muc, "nhieu_dieu", [source, target], None))
        print(f"{de_muc}: {quota} seeds ({len(multi)} multi, {len(deep)} deep-Khoan, {len(plain)} plain)")

    for number, seed in enumerate(seeds, start=1):
        seed["qid"] = f"ld{number:03d}"
    write_jsonl(SEEDS_PATH, [{"qid": s.pop("qid"), **s} for s in seeds])
    print(f"Wrote {len(seeds)} seeds to {SEEDS_PATH}")
    return 0


def make_seed(de_muc: str, style: str, mapcs: list[str], target: dict[str, Any] | None) -> dict[str, Any]:
    tags = []
    if target:
        tags.append("deep_khoan")
    if len(mapcs) > 1:
        tags.append("multi_dieu")
    return {
        "qid": "",
        "de_muc": de_muc,
        "style": style,
        "seed_mapcs": mapcs,
        "doc_types": [doc_type_of(m) for m in mapcs],
        "tags": tags,
        "target_khoan": target["khoan"] if target else None,
        "target_khoan_text": target["text"] if target else None,
    }


def parse_quotas(spec: str) -> dict[str, int]:
    quotas = {}
    for part in spec.split(","):
        key, _, value = part.partition("=")
        quotas[key.strip()] = int(value)
    return quotas


# ---------------------------------------------------------------------- draft

DRAFT_SYSTEM = (
    "Bạn soạn câu hỏi kiểm thử cho hệ thống tìm kiếm pháp luật lao động Việt Nam. "
    "Câu hỏi phải giống câu người dùng thật sẽ hỏi, và phải trả lời được bằng đoạn luật được cung cấp."
)


def draft_prompt(seed: dict[str, Any], chunks: dict[str, dict[str, Any]], feedback: str = "") -> str:
    style = seed["style"]
    if seed.get("target_khoan_text"):
        chunk = chunks[seed["seed_mapcs"][0]]
        passage = (
            f"Tiêu đề Điều: {chunk['ten']}\n"
            f"Khoản {seed['target_khoan']} của Điều này:\n{seed['target_khoan_text'][:2500]}\n\n"
            "Câu hỏi phải hỏi về nội dung RIÊNG của khoản này, không hỏi về các khoản khác."
        )
    else:
        limit = 1800 if len(seed["seed_mapcs"]) > 1 else 3000
        passage = "\n\n".join(
            f"Đoạn luật {i}: {chunks[m]['ten']}\n{article_body(chunks[m])[:limit]}"
            for i, m in enumerate(seed["seed_mapcs"], start=1)
        )
    return (
        f"{passage}\n\n"
        f"Yêu cầu kiểu câu hỏi: {STYLE_INSTRUCTIONS[style]}\n"
        "Quy tắc chung:\n"
        "- KHÔNG nhắc số Điều, số khoản, mã văn bản; KHÔNG chép nguyên tiêu đề Điều.\n"
        "- Câu hỏi phải trả lời được bằng đoạn luật trên, không cần kiến thức ngoài.\n"
        "- Tiếng Việt tự nhiên, có dấu.\n"
        f"{feedback}"
        "Trả về DUY NHẤT một JSON object:\n"
        '{"question": "...", "answer_hint": "<1 câu tóm tắt câu trả lời theo đoạn luật>", '
        f'"category": "<một trong: {", ".join(CATEGORIES)}>"}}'
    )


def check_draft(question: str, seed: dict[str, Any], chunks: dict[str, dict[str, Any]]) -> str:
    """Return a problem description, or '' when the draft is acceptable."""

    if len(question) < 8:
        return "câu hỏi quá ngắn"
    if ARTICLE_MENTION_RE.search(question):
        return "câu hỏi nhắc tới số Điều/mã văn bản"
    norm_question = normalize(question)
    for mapc in seed["seed_mapcs"]:
        title = normalize(chunks[mapc].get("ten", ""))
        if len(title) >= 25 and title in norm_question:
            return "câu hỏi chép nguyên tiêu đề Điều"
    return ""


def cmd_draft(args: argparse.Namespace) -> int:
    seeds = read_jsonl(SEEDS_PATH)
    if not seeds:
        raise SystemExit("Run `sample` first.")
    chunks = load_chunks(args.chunks)
    done = {qid for qid, row in latest_by_qid(read_jsonl(DRAFTS_PATH)).items() if row.get("status") == "ok"}
    todo = [s for s in seeds if s["qid"] not in done and (not args.qids or s["qid"] in args.qids)]
    todo = todo[: args.limit or None]
    print(f"{len(done)} drafted, {len(todo)} to draft")
    for index, seed in enumerate(todo, start=1):
        feedback, record = "", None
        for _attempt in range(3):
            try:
                reply = parse_json(llm_call(DRAFT_SYSTEM, draft_prompt(seed, chunks, feedback), max_tokens=800))
            except QuotaExhausted as exc:
                return stop_on_quota(exc, DRAFTS_PATH)
            except ValueError as exc:
                feedback = "- Lần trước trả về sai định dạng JSON, hãy trả về đúng JSON.\n"
                record = {"qid": seed["qid"], "status": "failed", "error": str(exc)[:200]}
                continue
            question = str(reply.get("question", "")).strip()
            problem = check_draft(question, seed, chunks)
            category = reply.get("category") if reply.get("category") in CATEGORIES else "khac"
            record = {
                "qid": seed["qid"],
                "status": "ok" if not problem else "failed",
                "drafted_by": "gemini",
                "question": question,
                "answer_hint": str(reply.get("answer_hint", "")).strip(),
                "category": category,
                "error": problem or None,
            }
            if not problem:
                break
            feedback = f"- Lần trước bị loại vì {problem}: \"{question}\". Hãy viết lại.\n"
        append_jsonl(DRAFTS_PATH, record)
        print(f"  [{index}/{len(todo)}] {seed['qid']} {seed['style']:<11} {record['status']}: {record.get('question', '')[:90]}")
        time.sleep(args.sleep)
    return 0


# ----------------------------------------------------------------------- pool


class BM25:
    def __init__(self, docs: list[list[str]], k1: float = 1.2, b: float = 0.75) -> None:
        self.k1, self.b = k1, b
        self.lengths = [len(d) for d in docs]
        self.avgdl = sum(self.lengths) / max(len(docs), 1)
        self.postings: dict[str, dict[int, int]] = defaultdict(dict)
        for index, doc in enumerate(docs):
            for term, tf in Counter(doc).items():
                self.postings[term][index] = tf
        n = len(docs)
        self.idf = {t: math.log(1 + (n - len(p) + 0.5) / (len(p) + 0.5)) for t, p in self.postings.items()}

    def top(self, query: list[str], k: int) -> list[int]:
        scores: dict[int, float] = defaultdict(float)
        for term in set(query):
            for index, tf in self.postings.get(term, {}).items():
                norm = tf + self.k1 * (1 - self.b + self.b * self.lengths[index] / self.avgdl)
                scores[index] += self.idf[term] * tf * (self.k1 + 1) / norm
        return sorted(scores, key=scores.get, reverse=True)[:k]


def corpus_embeddings(key: str, model: Any, texts: list[str], max_seq_length: int) -> Any:
    import numpy as np

    digest = hashlib.sha1(("\n\x00".join(texts) + f"|{max_seq_length}").encode("utf-8")).hexdigest()[:12]
    path = TESTSET_CACHE_DIR / f"{key}_{digest}.npy"
    if path.exists():
        return np.load(path)
    print(f"  encoding {len(texts)} Dieu with {key} (cached afterwards in {path.name})")
    vectors = model.encode(texts, batch_size=16, normalize_embeddings=True, show_progress_bar=True)
    path.parent.mkdir(parents=True, exist_ok=True)
    np.save(path, vectors)
    return vectors


def current_questions() -> list[dict[str, Any]]:
    """v1 questions + successful drafts, with text edited in review.xlsx taking precedence."""

    seeds = {s["qid"]: s for s in read_jsonl(SEEDS_PATH)}
    drafts = latest_by_qid(read_jsonl(DRAFTS_PATH))
    questions: list[dict[str, Any]] = []
    for row in read_jsonl(V1_QUESTIONS):
        questions.append(
            {
                "qid": row["id"],
                "question": row["question"],
                "question_draft": row["question"],
                "de_muc": "20.5",
                "style": "truc_tiep",
                "tags": [],
                "seed_mapcs": [m for m, g in row["relevance"].items() if float(g) >= 3] or list(row["relevance"]),
                "target_khoan": None,
                "answer_hint": "",
                "category": row.get("category", "khac"),
                "origin": "manual_v1",
                "drafted_by": "human",
                "v1_relevance": {m: int(g) for m, g in row["relevance"].items()},
                "notes": row.get("notes", ""),
            }
        )
    for qid, seed in seeds.items():
        draft = drafts.get(qid)
        if not draft or draft.get("status") != "ok":
            continue
        questions.append(
            {
                "qid": qid,
                "question": draft["question"],
                "question_draft": draft["question"],
                "de_muc": seed["de_muc"],
                "style": seed["style"],
                "tags": seed["tags"],
                "seed_mapcs": seed["seed_mapcs"],
                "target_khoan": seed.get("target_khoan"),
                "answer_hint": draft.get("answer_hint", ""),
                "category": draft.get("category", "khac"),
                "origin": "llm_draft",
                "drafted_by": draft.get("drafted_by", "gemini"),
                "v1_relevance": {},
                "notes": "",
            }
        )
    edited = read_review_questions() if REVIEW_PATH.exists() else {}
    for question in questions:
        row = edited.get(question["qid"])
        if row and (row.get("question") or "").strip():
            question["question"] = str(row["question"]).strip()
    return questions


def cmd_pool(args: argparse.Namespace) -> int:
    import numpy as np
    from sentence_transformers import SentenceTransformer

    chunks = load_chunks(args.chunks)
    mapcs = list(chunks)
    texts = [chunks[m]["text"] for m in mapcs]
    questions = current_questions()
    if not questions:
        raise SystemExit("No questions yet; run `draft` first.")
    query_texts = [q["question"] for q in questions]
    rankings: dict[str, list[list[int]]] = {}

    bm25 = BM25([tokens(t) for t in texts])
    rankings["bm25"] = [bm25.top(tokens(q), args.depth) for q in query_texts]
    for key in args.models:
        print(f"[{key}] {DENSE_MODELS[key]}")
        model = SentenceTransformer(DENSE_MODELS[key])
        model.max_seq_length = min(model.max_seq_length or args.max_seq_length, args.max_seq_length)
        corpus = corpus_embeddings(key, model, texts, model.max_seq_length)
        queries = model.encode(query_texts, batch_size=32, normalize_embeddings=True)
        scores = np.asarray(queries) @ np.asarray(corpus).T
        rankings[key] = [list(np.argsort(-row)[: args.depth]) for row in scores]
        del model

    pools = []
    for qi, question in enumerate(questions):
        candidates: dict[str, dict[str, Any]] = {}
        for system, per_query in rankings.items():
            for rank, doc_index in enumerate(per_query[qi], start=1):
                entry = candidates.setdefault(mapcs[int(doc_index)], {"systems": {}})
                entry["systems"][system] = rank
        for mapc in [*question["seed_mapcs"], *question["v1_relevance"]]:
            if mapc in chunks:
                candidates.setdefault(mapc, {"systems": {}})
        for mapc, entry in candidates.items():
            entry["is_seed"] = mapc in question["seed_mapcs"]
        pools.append({"qid": question["qid"], "question": question["question"], "candidates": candidates})
    write_jsonl(POOLS_PATH, pools)
    sizes = [len(p["candidates"]) for p in pools]
    found = sum(any(c["systems"] for m, c in p["candidates"].items() if c["is_seed"]) for p in pools)
    print(f"Wrote {len(pools)} pools to {POOLS_PATH} (avg {sum(sizes) / len(sizes):.1f} candidates; "
          f"seed retrieved by some system for {found}/{len(pools)})")
    return 0


# ---------------------------------------------------------------------- judge

JUDGE_SYSTEM = (
    "Bạn là luật sư lao động Việt Nam, chấm mức độ liên quan giữa câu hỏi và các điều luật "
    "để xây dựng bộ đánh giá hệ thống tìm kiếm. Chấm khách quan, chặt chẽ."
)
GRADE_RUBRIC = (
    "Thang điểm:\n"
    "3 = điều luật trả lời trực tiếp, đầy đủ câu hỏi (hoặc phần chính của câu hỏi).\n"
    "2 = chứa một phần quan trọng của câu trả lời, cần kết hợp với điều khác.\n"
    "1 = có liên quan/cung cấp ngữ cảnh, nhưng không trả lời câu hỏi.\n"
    "0 = không liên quan.\n"
)


def candidate_snippet(chunk: dict[str, Any], question: str, limit: int) -> str:
    """Title + body; for long Dieu, the opening plus the Khoan closest to the question."""

    body = article_body(chunk)
    if len(body) <= limit:
        return body
    khoan = split_khoan(body)
    if not khoan:
        return body[:limit] + " …"
    query_terms = set(tokens(question))
    best = sorted(khoan, key=lambda k: len(query_terms & set(tokens(k["text"]))), reverse=True)[:2]
    best.sort(key=lambda k: k["khoan"])
    head = body[: min(300, khoan[0]["start"] or 300)]
    parts = [head.strip(), *[k["text"] for k in best]]
    return " … ".join(p for p in parts if p)[:limit] + " …"


def cmd_judge(args: argparse.Namespace) -> int:
    chunks = load_chunks(args.chunks)
    pools = read_jsonl(POOLS_PATH)
    if not pools:
        raise SystemExit("Run `pool` first.")
    judged = latest_by_qid(read_jsonl(JUDGMENTS_PATH))
    todo = []
    for pool in pools:
        previous = judged.get(pool["qid"])
        known = previous["grades"] if previous and previous.get("question") == pool["question"] else {}
        missing = [m for m in pool["candidates"] if m not in known]
        if missing:
            todo.append((pool, known, missing))
    todo = todo[: args.limit or None]
    print(f"{len(pools) - len(todo)} pools fully judged, {len(todo)} to judge")
    for index, (pool, known, missing) in enumerate(todo, start=1):
        listing = "\n\n".join(
            f"[{i}] {chunks[m]['ten']}\n{candidate_snippet(chunks[m], pool['question'], args.snippet_chars)}"
            for i, m in enumerate(missing)
        )
        user = (
            f"Câu hỏi: {pool['question']}\n\n{GRADE_RUBRIC}\nCác điều luật ứng viên:\n{listing}\n\n"
            'Trả về DUY NHẤT JSON: {"grades": [{"i": <chỉ số>, "grade": <0-3>, "why": "<tối đa 15 từ>"}, ...]} '
            "cho TẤT CẢ ứng viên."
        )
        grades = dict(known)
        try:
            raw = llm_call(JUDGE_SYSTEM, user, max_tokens=400 + 80 * len(missing))
        except QuotaExhausted as exc:
            return stop_on_quota(exc, JUDGMENTS_PATH)
        try:
            reply = parse_json(raw)
            items = reply.get("grades", []) if isinstance(reply, dict) else []
        except ValueError:
            items = salvage_grades(raw)  # truncated reply: keep the complete items
        for item in items:
            i, grade = item.get("i"), item.get("grade")
            if isinstance(i, int) and 0 <= i < len(missing) and grade in (0, 1, 2, 3):
                grades[missing[i]] = {"grade": grade, "why": str(item.get("why", ""))[:120]}
        append_jsonl(JUDGMENTS_PATH, {"qid": pool["qid"], "question": pool["question"], "grades": grades})
        print(f"  [{index}/{len(todo)}] {pool['qid']}: {len(grades)}/{len(pool['candidates'])} graded")
        time.sleep(args.sleep)
    return 0


# --------------------------------------------------------------- review.xlsx

Q_COLUMNS = [
    ("qid", 8), ("de_muc", 7), ("style", 11), ("tags", 12), ("origin", 10), ("drafted_by", 10), ("seed_mapcs", 18),
    ("seed_titles", 40), ("target_khoan", 6), ("answer_hint", 40), ("question_draft", 45),
    ("question", 55), ("category", 22), ("action", 8), ("notes", 30),
]
J_COLUMNS = [
    ("qid", 8), ("question", 45), ("mapc", 16), ("ten", 40), ("is_seed", 6), ("pooled_by", 22),
    ("suggested_grade", 8), ("suggested_reason", 30), ("grade", 7), ("snippet", 80), ("full_text", 30),
]
Q_HUMAN = ("question", "category", "action", "notes")
J_HUMAN = ("grade",)


def _read_sheet(path: Path, sheet: str) -> list[dict[str, Any]]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    rows = list(workbook[sheet].iter_rows(values_only=True))
    workbook.close()
    if not rows:
        return []
    header = [str(h) for h in rows[0]]
    return [dict(zip(header, row)) for row in rows[1:] if any(v not in (None, "") for v in row)]


def read_review_questions() -> dict[str, dict[str, Any]]:
    return {str(r["qid"]): r for r in _read_sheet(REVIEW_PATH, "questions")}


def read_review_judgments() -> dict[tuple[str, str], dict[str, Any]]:
    return {(str(r["qid"]), str(r["mapc"])): r for r in _read_sheet(REVIEW_PATH, "judgments")}


def cmd_export(args: argparse.Namespace) -> int:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    chunks = load_chunks(args.chunks)
    questions = current_questions()
    pools = {p["qid"]: p for p in read_jsonl(POOLS_PATH)}
    judged = latest_by_qid(read_jsonl(JUDGMENTS_PATH))
    old_q = read_review_questions() if REVIEW_PATH.exists() else {}
    old_j = read_review_judgments() if REVIEW_PATH.exists() else {}

    workbook = Workbook()
    q_sheet = workbook.active
    q_sheet.title = "questions"
    j_sheet = workbook.create_sheet("judgments")
    human_fill = PatternFill("solid", fgColor="FFF2CC")

    def setup(sheet: Any, columns: list[tuple[str, int]], human: tuple[str, ...]) -> None:
        sheet.append([name for name, _ in columns])
        for index, (name, width) in enumerate(columns, start=1):
            letter = sheet.cell(row=1, column=index).column_letter
            sheet.column_dimensions[letter].width = width
            sheet.cell(row=1, column=index).font = Font(bold=True)
            if name in human:
                sheet.cell(row=1, column=index).fill = human_fill
        sheet.freeze_panes = "B2"

    setup(q_sheet, Q_COLUMNS, Q_HUMAN)
    setup(j_sheet, J_COLUMNS, J_HUMAN)

    n_judgment_rows = 0
    for question in questions:
        qid = question["qid"]
        previous = old_q.get(qid, {})
        default_action = "keep" if question["origin"] == "manual_v1" else ""
        q_row = {
            "qid": qid,
            "de_muc": question["de_muc"],
            "style": question["style"],
            "tags": ",".join(question["tags"]),
            "origin": question["origin"],
            "drafted_by": question["drafted_by"],
            "seed_mapcs": ", ".join(question["seed_mapcs"]),
            "seed_titles": " | ".join(chunks[m]["ten"] for m in question["seed_mapcs"] if m in chunks),
            "target_khoan": question["target_khoan"],
            "answer_hint": question["answer_hint"],
            "question_draft": question["question_draft"],
            "question": previous.get("question") or question["question"],
            "category": previous.get("category") or question["category"],
            "action": previous.get("action") if previous.get("action") not in (None, "") else default_action,
            "notes": previous.get("notes") or question["notes"],
        }
        q_sheet.append([q_row[name] for name, _ in Q_COLUMNS])

        pool = pools.get(qid)
        if not pool:
            continue
        suggestions = judged.get(qid, {}).get("grades", {})
        candidates = pool["candidates"]

        def order(mapc: str) -> tuple:
            entry = candidates[mapc]
            suggested = suggestions.get(mapc, {}).get("grade", -1)
            return (not entry["is_seed"], -suggested, min(entry["systems"].values(), default=99))

        for mapc in sorted(candidates, key=order):
            entry, suggestion = candidates[mapc], suggestions.get(mapc, {})
            chunk = chunks[mapc]
            previous_grade = old_j.get((qid, mapc), {}).get("grade")
            grade = previous_grade if previous_grade not in (None, "") else question["v1_relevance"].get(mapc)
            j_row = {
                "qid": qid,
                "question": q_row["question"],
                "mapc": mapc,
                "ten": chunk["ten"],
                "is_seed": "x" if entry["is_seed"] else "",
                "pooled_by": ", ".join(f"{s}#{r}" for s, r in sorted(entry["systems"].items(), key=lambda kv: kv[1])),
                "suggested_grade": suggestion.get("grade"),
                "suggested_reason": suggestion.get("why", ""),
                "grade": grade,
                "snippet": candidate_snippet(chunk, q_row["question"], 700),
                "full_text": article_body(chunk)[:32000],
            }
            j_sheet.append([j_row[name] for name, _ in J_COLUMNS])
            n_judgment_rows += 1

    wrap = Alignment(wrap_text=True, vertical="top")
    for sheet, columns in ((q_sheet, Q_COLUMNS), (j_sheet, J_COLUMNS)):
        for name in ("seed_titles", "answer_hint", "question_draft", "question", "ten", "snippet", "suggested_reason"):
            if name in [c for c, _ in columns]:
                letter = sheet.cell(row=1, column=[c for c, _ in columns].index(name) + 1).column_letter
                for cell in sheet[letter][1:]:
                    cell.alignment = wrap
    q_last, j_last = q_sheet.max_row, j_sheet.max_row

    def column_range(columns: list[tuple[str, int]], name: str, last_row: int) -> str:
        from openpyxl.utils import get_column_letter

        letter = get_column_letter([c for c, _ in columns].index(name) + 1)
        return f"{letter}2:{letter}{max(last_row, 2)}"

    action_rule = DataValidation(type="list", formula1='"keep,drop"', allow_blank=True)
    action_rule.add(column_range(Q_COLUMNS, "action", q_last))
    # Inline list formulas are capped at 255 characters, so categories live on a hidden sheet.
    lists = workbook.create_sheet("lists")
    for category in CATEGORIES:
        lists.append([category])
    lists.sheet_state = "hidden"
    category_rule = DataValidation(type="list", formula1=f"=lists!$A$1:$A${len(CATEGORIES)}", allow_blank=True)
    category_rule.add(column_range(Q_COLUMNS, "category", q_last))
    grade_rule = DataValidation(type="list", formula1='"0,1,2,3"', allow_blank=True)
    grade_rule.add(column_range(J_COLUMNS, "grade", j_last))
    q_sheet.add_data_validation(action_rule)
    q_sheet.add_data_validation(category_rule)
    j_sheet.add_data_validation(grade_rule)
    q_sheet.auto_filter.ref = q_sheet.dimensions
    j_sheet.auto_filter.ref = j_sheet.dimensions

    REVIEW_PATH.parent.mkdir(parents=True, exist_ok=True)
    try:
        workbook.save(REVIEW_PATH)
    except PermissionError:
        raise SystemExit(f"Cannot write {REVIEW_PATH}: close it in Excel first.")
    print(f"Wrote {REVIEW_PATH}: {len(questions)} questions, {n_judgment_rows} judgment rows "
          f"(kept {len(old_q)} reviewed question rows / {sum(1 for r in old_j.values() if r.get('grade') not in (None, ''))} grades)")
    return 0


# ------------------------------------------------------------------- finalize


def _action(value: Any) -> str:
    text = str(value or "").strip().lower()
    if text in ("keep", "k", "x", "1", "yes", "y", "giu", "giữ"):
        return "keep"
    if text in ("drop", "d", "0", "no", "n", "bo", "bỏ"):
        return "drop"
    return ""


def _grade(value: Any) -> int | None:
    if value in (None, ""):
        return None
    try:
        grade = int(float(value))
    except (TypeError, ValueError):
        return None
    return grade if grade in (0, 1, 2, 3) else None


def cmd_finalize(args: argparse.Namespace) -> int:
    if not REVIEW_PATH.exists():
        raise SystemExit("Run `export` and review the workbook first.")
    chunks = load_chunks(args.chunks)
    meta = {q["qid"]: q for q in current_questions()}
    reviewed_q = read_review_questions()
    judgments: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in _read_sheet(REVIEW_PATH, "judgments"):
        judgments[str(row["qid"])].append(row)

    errors: list[str] = []
    warnings: list[str] = []
    output: list[dict[str, Any]] = []
    counts = Counter()
    for qid, row in reviewed_q.items():
        action = _action(row.get("action"))
        if action == "drop":
            counts["dropped"] += 1
            continue
        if not action:
            counts["unreviewed"] += 1
            if not args.allow_unreviewed:
                errors.append(f"{qid}: action is empty (keep/drop)")
            continue
        question_text = str(row.get("question") or "").strip()
        relevance: dict[str, int] = {}
        missing, from_suggestion = [], 0
        for judgment in judgments.get(qid, []):
            grade = _grade(judgment.get("grade"))
            if grade is None and args.use_suggested:
                grade = _grade(judgment.get("suggested_grade"))
                from_suggestion += grade is not None
            if grade is None:
                missing.append(str(judgment["mapc"]))
                continue
            if grade > 0:
                relevance[str(judgment["mapc"])] = grade
            if str(judgment.get("is_seed")) == "x" and grade == 0:
                warnings.append(f"{qid}: seed {judgment['mapc']} graded 0 -- does the question still match its seed?")
        if missing:
            errors.append(f"{qid}: {len(missing)} candidate(s) without grade, e.g. {missing[0]}")
            continue
        if not any(grade >= 2 for grade in relevance.values()):
            errors.append(f"{qid}: no Dieu graded 2 or 3 -- drop the question or fix the grades")
            continue
        unknown = [m for m in relevance if m not in chunks]
        if unknown:
            errors.append(f"{qid}: unknown mapc {unknown[0]}")
            continue
        info = meta.get(qid, {})
        edited = question_text != str(info.get("question_draft", question_text)).strip()
        origin = info.get("origin", "llm_draft")
        if origin == "llm_draft":
            origin = "llm_draft_edited" if edited else "llm_draft_reviewed"
        output.append(
            {
                "id": qid,
                "question": question_text,
                "de_muc": info.get("de_muc", row.get("de_muc")),
                "category": row.get("category") or info.get("category", "khac"),
                "style": info.get("style", row.get("style")),
                "tags": info.get("tags", []),
                "relevance": dict(sorted(relevance.items(), key=lambda kv: -kv[1])),
                "seed_mapcs": info.get("seed_mapcs", []),
                "target_khoan": info.get("target_khoan"),
                "origin": origin,
                "drafted_by": info.get("drafted_by", "unknown"),
                "grades_from_suggestion": from_suggestion,
                "notes": row.get("notes") or "",
            }
        )

    seen: dict[str, str] = {}
    for item in output:
        key = normalize(item["question"])
        if key in seen:
            warnings.append(f"{item['id']}: same question as {seen[key]}")
        seen[key] = item["id"]

    for message in warnings:
        print(f"WARNING {message}")
    for message in errors[:50]:
        print(f"ERROR   {message}")
    print(f"{len(output)} kept, {counts['dropped']} dropped, {counts['unreviewed']} unreviewed, "
          f"{len(errors)} error(s), {len(warnings)} warning(s)")
    if errors:
        print("Nothing written. Fix the workbook, or see --allow-unreviewed / --use-suggested.")
        return 1
    output.sort(key=lambda q: (not q["id"].startswith("q"), q["id"]))
    write_jsonl(args.output, output)
    print(f"Wrote {args.output}")

    from labor_rag.evaluation.embedding_benchmark import load_and_validate_dataset

    load_and_validate_dataset(args.chunks, args.output)
    print("Benchmark validation passed.")
    return 0


# ---------------------------------------------------------------------- stats


def cmd_stats(args: argparse.Namespace) -> int:
    chunks = load_chunks(args.chunks)
    path = args.output if args.output.exists() else None
    if path:
        questions = read_jsonl(path)
        source = str(path)
        rows = [
            {"de_muc": q["de_muc"], "style": q["style"], "tags": q.get("tags", []), "relevance": q["relevance"]}
            for q in questions
        ]
    else:  # before finalize: plan coverage from seeds + v1
        source = f"{SEEDS_PATH} + {V1_QUESTIONS} (not finalized yet)"
        rows = [
            {"de_muc": "20.5", "style": "truc_tiep", "tags": [], "relevance": q["relevance"]}
            for q in read_jsonl(V1_QUESTIONS)
        ] + [
            {"de_muc": s["de_muc"], "style": s["style"], "tags": s["tags"], "relevance": {m: 3 for m in s["seed_mapcs"]}}
            for s in read_jsonl(SEEDS_PATH)
        ]
    gold = [m for r in rows for m in r["relevance"]]
    eligible = Counter(de_muc_of(m) for m, c in chunks.items() if is_eligible(c))
    covered = Counter(de_muc_of(m) for m in set(gold))
    stats = {
        "source": source,
        "questions": len(rows),
        "by_de_muc": dict(sorted(Counter(r["de_muc"] for r in rows).items())),
        "by_style": dict(Counter(r["style"] for r in rows).most_common()),
        "by_tag": dict(Counter(t for r in rows for t in r["tags"]).most_common()),
        "gold_by_doc_type": dict(Counter(doc_type_of(m) for m in gold).most_common()),
        "gold_per_question": round(len(gold) / max(len(rows), 1), 2),
        "grade_distribution": dict(sorted(Counter(int(g) for r in rows for g in r["relevance"].values()).items())),
        "dieu_coverage_by_de_muc": {
            d: f"{covered.get(d, 0)}/{n} ({covered.get(d, 0) / n:.0%})" for d, n in sorted(eligible.items())
        },
    }
    STATS_PATH.parent.mkdir(parents=True, exist_ok=True)
    STATS_PATH.write_text(json.dumps(stats, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(stats, ensure_ascii=False, indent=2))
    return 0


# ------------------------------------------------------------------------ cli


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--chunks", type=Path, default=DEFAULT_CHUNKS)
    sub = parser.add_subparsers(dest="command", required=True)

    sample = sub.add_parser("sample", help="choose seed Dieu")
    sample.add_argument("--edges", type=Path, default=DEFAULT_EDGES)
    sample.add_argument("--quota", help='e.g. "20.2=50,20.3=20,20.4=30,20.5=20"')
    sample.add_argument("--seed", type=int, default=42)
    sample.add_argument("--force", action="store_true")
    sample.set_defaults(func=cmd_sample)

    draft = sub.add_parser("draft", help="LLM drafts one question per seed")
    draft.add_argument("--limit", type=int, default=0, help="only draft this many (0 = all)")
    draft.add_argument("--sleep", type=float, default=4.0, help="seconds between calls (free-tier RPM)")
    draft.add_argument("--qids", nargs="+", help="only these seeds (e.g. a trial run)")
    draft.set_defaults(func=cmd_draft)

    pool = sub.add_parser("pool", help="candidate Dieu from BM25 + embedding models")
    pool.add_argument("--models", nargs="+", default=list(DENSE_MODELS), choices=list(DENSE_MODELS))
    pool.add_argument("--depth", type=int, default=5, help="top-k taken from each system")
    pool.add_argument("--max-seq-length", type=int, default=512)
    pool.set_defaults(func=cmd_pool)

    judge = sub.add_parser("judge", help="LLM suggests a 0-3 grade per candidate")
    judge.add_argument("--limit", type=int, default=0)
    judge.add_argument("--sleep", type=float, default=4.0)
    judge.add_argument("--snippet-chars", type=int, default=1500)
    judge.set_defaults(func=cmd_judge)

    export = sub.add_parser("export", help="write/refresh review.xlsx, keeping reviewer input")
    export.set_defaults(func=cmd_export)

    finalize = sub.add_parser("finalize", help="review.xlsx -> benchmark_questions_v2.jsonl")
    finalize.add_argument("--output", type=Path, default=OUTPUT_QUESTIONS)
    finalize.add_argument("--allow-unreviewed", action="store_true", help="skip questions with an empty action")
    finalize.add_argument("--use-suggested", action="store_true",
                          help="fill empty grades with the LLM suggestion (counted per question; avoid for the final set)")
    finalize.set_defaults(func=cmd_finalize)

    stats = sub.add_parser("stats", help="coverage report")
    stats.add_argument("--output", type=Path, default=OUTPUT_QUESTIONS)
    stats.set_defaults(func=cmd_stats)
    return parser


def main(argv: list[str] | None = None) -> int:
    sys.stdout.reconfigure(encoding="utf-8")
    args = build_parser().parse_args(argv)
    return int(args.func(args) or 0)


if __name__ == "__main__":
    raise SystemExit(main())
